from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, RGBColor, Inches

NAVY=RGBColor(0x1B,0x3A,0x5C); GREEN=RGBColor(0x1E,0x9B,0x6E); GRAY=RGBColor(0x33,0x33,0x33)
BLOCKS={"h1","h2","h3","p","li","hr"}

class P(HTMLParser):
    def __init__(self):
        super().__init__(); self.out=[]; self.cur=None; self.b=0; self.i=0; self.q=0
    def handle_starttag(self,t,a):
        if t=="blockquote": self.q+=1
        elif t=="strong": self.b+=1
        elif t=="em": self.i+=1
        elif t=="hr": self.out.append(("hr",[],self.q))
        elif t in BLOCKS: self.cur=(t,[],self.q)
    def handle_endtag(self,t):
        if t=="blockquote": self.q=max(0,self.q-1)
        elif t=="strong": self.b=max(0,self.b-1)
        elif t=="em": self.i=max(0,self.i-1)
        elif t in BLOCKS and self.cur:
            if any(x.strip() for x,_,_ in self.cur[1]): self.out.append(self.cur)
            self.cur=None
    def handle_data(self,d):
        if self.cur and d: self.cur[1].append((d,self.b>0,self.i>0))

p=P(); p.feed(open("outreach.html").read())
doc=Document()
st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(11); st.font.color.rgb=GRAY
doc.sections[0].left_margin=doc.sections[0].right_margin=Inches(1)

def emit(par,spans,color=None):
    for text,bold,ital in spans:
        t=" ".join(text.split())
        if not t: continue
        if par.runs and not par.runs[-1].text.endswith(" "): t=" "+t
        r=par.add_run(t); r.bold=bold; r.italic=ital
        if color: r.font.color.rgb=color

for tag,spans,q in p.out:
    if tag=="hr": doc.add_paragraph(); continue
    if q:
        par=doc.add_paragraph()
        par.paragraph_format.left_indent=Inches(0.35)
        par.paragraph_format.space_after=Pt(8)
        emit(par,spans,RGBColor(0,0,0)); continue
    if tag=="h1":
        par=doc.add_paragraph(); emit(par,spans,NAVY)
        for r in par.runs: r.bold=True; r.font.size=Pt(22)
    elif tag=="h2":
        doc.add_paragraph(); par=doc.add_paragraph(); emit(par,spans,NAVY)
        for r in par.runs: r.bold=True; r.font.size=Pt(15)
    elif tag=="h3":
        par=doc.add_paragraph(); emit(par,spans,GREEN)
        for r in par.runs: r.bold=True; r.font.size=Pt(12)
    elif tag=="li":
        par=doc.add_paragraph(style="List Bullet"); emit(par,spans)
    else:
        par=doc.add_paragraph(); emit(par,spans)

out="/home/user/BaxterMortgage/content-marketing/Outreach_Playbook_Derek_Smith.docx"
doc.save(out); print("saved",out)
